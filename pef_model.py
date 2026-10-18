"""
Motor de cálculo del Plan Económico Financiero (PEF) de concesión vial.

Replica las fórmulas del modelo Excel PEF_X (hojas CONTROL, INPUTS, HIPOTESIS,
CAPEX, OPEX, FLUJO, PRESTAMO, IMPUESTOS y AMORTIZACIONES), incluidas sus
particularidades, para obtener los mismos resultados que la planilla.

Uso:
    from pef_model import cargar_inputs_excel, calcular
    inp = cargar_inputs_excel("PEF_X.xlsx")
    res = calcular(inp)
    print(res.kpis)
"""
from __future__ import annotations

import copy
import math
from dataclasses import dataclass, field, asdict
from typing import Optional

import numpy as np
import pandas as pd

N = 20  # años de la grilla del modelo


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------
def _unos():
    return [1.0] * N


def _ceros():
    return [0.0] * N


@dataclass
class Inputs:
    nombre: str = "PEF"
    # Escenario y fechas
    escenario: int = 2                 # 1 = 15 años, 2 = 20 años   (CONTROL C12)
    anio_inicio: int = 2027            # año calendario del año 1   (CONTROL C15)
    # Tránsito y tarifa
    uteq: float = 7_200_000            # UTEQUIS base               (CONTROL C20)
    uteq_adicional: float = 0.0        # UTEQUIS peajes adicionales (CONTROL C21)
    uteq_base_desde_anio3: Optional[float] = None  # base de tránsito distinta desde el año 3 (FLUJO D114, caso PEF_VI)
    tarifa: float = 4_680              # $ por UTEQUI               (CONTROL C24)
    crecimiento: float = 0.03          # crecimiento anual          (CONTROL C36)
    nuevos_peajes: list = field(default_factory=lambda: [1.0, 1.0, 1.0, 1.0])  # años 2..5 (CONTROL C37:C40)
    otros_ingresos_pct: float = 0.01   # % de peajes (FLUJO fila 75)
    # Sensibilidades (CONTROL C29:C35)
    sens_obras: float = 0.0
    sens_repav: float = 0.0
    sens_cym: float = 0.0
    sens_atu: float = 0.0
    sens_adm: float = 0.0
    sens_expl: float = 0.0
    # Financiamiento
    tasa_descuento: float = 0.10       # CONTROL G21
    pct_deuda: float = 0.10            # CONTROL G25 (sobre CAPEX total)
    tna: float = 0.085                 # CONTROL G26
    cuotas: int = 6                    # CONTROL G28
    sistema: int = 1                   # 1 francés, 2 alemán (CONTROL G29)
    gracia: bool = True                # CONTROL G32 ("si"/"NO")
    gastos_prestamo_pct: float = 0.01  # PRESTAMO C10
    # Impuestos (CONTROL C43:C56)
    iibb: float = 0.025
    ganancias: float = 0.35
    idc: float = 0.012
    iva: float = 0.21
    iva_cym: float = 0.15
    iva_atu: float = 0.08
    iva_adm: float = 0.057
    iva_expl: float = 0.18
    iva_capex: float = 0.21
    sellos: float = 0.012
    municipales: float = 0.005
    pct_amortizable: float = 1.0
    pct_deduccion_intereses: float = 0.30
    # Amortización impositiva
    variante_amort: int = 1            # 1: obras plazo restante + repav vida útil; 2: todo a 5 años; 3: instantánea (CONTROL L16)
    vida_util_repav: int = 5           # CONTROL C57
    vida_util_obras_5: int = 5         # CONTROL C58
    # Garantías (CONTROL E36:G56)
    gar_pr_contrato: float = 0.0       # F39 monto garantía a primer requerimiento (contrato)
    gar_pr_obras_1a5: float = 0.0      # F40
    gar_pr_obras_6a20: float = 0.0     # F41
    gar_pr_costo_contrato: float = 0.01  # G39
    gar_pr_costo_obras: float = 0.01     # G40
    caucion_pct_obras: float = 0.03      # F50 = ROUND(obras+repav * 3%, -8)
    caucion_pct_contrato: float = 0.015  # F51 = ROUND(OPEX * 1,5%, -8)
    caucion_costo_obras: float = 0.01    # G50
    caucion_costo_contrato: float = 0.01 # G51
    # Series anuales (INPUTS), montos con IVA en $
    obras: list = field(default_factory=_ceros)
    obras_f1: list = field(default_factory=_unos)
    obras_f2: list = field(default_factory=_unos)
    repav: list = field(default_factory=_ceros)
    repav_f1: list = field(default_factory=_unos)
    repav_f2: list = field(default_factory=_unos)
    operacion: list = field(default_factory=_ceros)
    operacion_f1: list = field(default_factory=_unos)
    operacion_f2: list = field(default_factory=_unos)
    mantenimiento: list = field(default_factory=_ceros)
    mantenimiento_f1: list = field(default_factory=_unos)
    mantenimiento_f2: list = field(default_factory=_unos)
    servicios: list = field(default_factory=_ceros)
    servicios_f1: list = field(default_factory=_unos)
    servicios_f2: list = field(default_factory=_unos)
    administracion: list = field(default_factory=_ceros)
    administracion_f1: list = field(default_factory=_unos)
    administracion_f2: list = field(default_factory=_unos)
    # Factores HIPOTESIS
    h_peajes: list = field(default_factory=_unos)       # fila 18
    h_peajes_nuevos: list = field(default_factory=_unos)  # fila 19
    h_otros: list = field(default_factory=_unos)        # fila 20
    h_cym: list = field(default_factory=_unos)          # fila 24
    h_atu: list = field(default_factory=_unos)          # fila 25
    h_adm: list = field(default_factory=_unos)          # fila 26
    h_expl: list = field(default_factory=_unos)         # fila 27
    h_pv: list = field(default_factory=_unos)           # fila 31
    h_oo: list = field(default_factory=_unos)           # fila 32
    h_rep: list = field(default_factory=_unos)          # fila 33

    @property
    def plazo(self) -> int:
        return 20 if self.escenario == 2 else 15

    def copia(self, **cambios) -> "Inputs":
        nuevo = copy.deepcopy(self)
        for k, v in cambios.items():
            setattr(nuevo, k, v)
        return nuevo

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "Inputs":
        campos = {f for f in cls.__dataclass_fields__}
        return cls(**{k: v for k, v in d.items() if k in campos})


# ---------------------------------------------------------------------------
# Lectura desde el Excel
# ---------------------------------------------------------------------------
def _num(v, defecto=0.0):
    if v is None or isinstance(v, str):
        return defecto
    try:
        return float(v)
    except (TypeError, ValueError):
        return defecto


def cargar_inputs_excel(archivo, nombre: str | None = None) -> Inputs:
    """Lee los datos de entrada de un PEF con estructura CONTROL/FLUJO/INPUTS.

    `archivo` puede ser una ruta o un objeto tipo archivo (por ejemplo, el de
    st.file_uploader). Se leen los valores guardados por Excel.
    """
    import openpyxl

    wb = openpyxl.load_workbook(archivo, data_only=True)
    C = wb["CONTROL"]
    I = wb["INPUTS"]
    H = wb["HIPOTESIS"]
    F = wb["FLUJO"]
    P = wb["PRESTAMO "] if "PRESTAMO " in wb.sheetnames else wb["PRESTAMO"]

    fila = lambda ws, r, c0: [_num(ws.cell(r, c0 + i).value, 0.0) for i in range(N)]
    fila1 = lambda ws, r, c0: [_num(ws.cell(r, c0 + i).value, 1.0) for i in range(N)]

    fecha = C["C15"].value
    anio = getattr(fecha, "year", 2027)
    gracia_txt = str(C["G32"].value or "si").strip().upper()
    d114 = F["D114"].value

    inp = Inputs(
        nombre=nombre or "PEF",
        escenario=int(_num(C["C12"].value, 2)),
        anio_inicio=anio,
        uteq=_num(C["C20"].value),
        uteq_adicional=_num(C["C21"].value),
        uteq_base_desde_anio3=_num(d114) if isinstance(d114, (int, float)) else None,
        tarifa=_num(C["C24"].value),
        crecimiento=_num(C["C36"].value, 0.03),
        nuevos_peajes=[_num(C[f"C{r}"].value, 1.0) for r in (37, 38, 39, 40)],
        sens_obras=_num(C["C29"].value),
        sens_repav=_num(C["C30"].value),
        sens_cym=_num(C["C32"].value),
        sens_atu=_num(C["C33"].value),
        sens_adm=_num(C["C34"].value),
        sens_expl=_num(C["C35"].value),
        tasa_descuento=_num(C["G21"].value, 0.10),
        pct_deuda=_num(C["G25"].value),
        tna=_num(C["G26"].value),
        cuotas=int(_num(C["G28"].value, 6)),
        sistema=int(_num(C["G29"].value, 1)),
        gracia=(gracia_txt != "NO"),
        gastos_prestamo_pct=_num(P["C10"].value, 0.01),
        iibb=_num(C["C43"].value),
        ganancias=_num(C["C44"].value),
        idc=_num(C["C45"].value),
        iva=_num(C["C46"].value, 0.21),
        iva_cym=_num(C["C48"].value),
        iva_atu=_num(C["C49"].value),
        iva_adm=_num(C["C50"].value),
        iva_expl=_num(C["C51"].value),
        iva_capex=_num(C["C52"].value, 0.21),
        sellos=_num(C["C53"].value),
        municipales=_num(C["C54"].value),
        pct_amortizable=_num(C["C55"].value, 1.0),
        pct_deduccion_intereses=_num(C["C56"].value, 0.30),
        variante_amort=int(_num(C["L16"].value, 1)),
        vida_util_repav=int(_num(C["C57"].value, 5)) or 5,
        vida_util_obras_5=int(_num(C["C58"].value, 5)) or 5,
        gar_pr_contrato=_num(C["F39"].value),
        gar_pr_obras_1a5=_num(C["F40"].value),
        gar_pr_obras_6a20=_num(C["F41"].value),
        gar_pr_costo_contrato=_num(C["G39"].value, 0.01),
        gar_pr_costo_obras=_num(C["G40"].value, 0.01),
        caucion_costo_obras=_num(C["G50"].value, 0.01),
        caucion_costo_contrato=_num(C["G51"].value, 0.01),
        obras=fila(I, 4, 2), obras_f1=fila1(I, 5, 2), obras_f2=fila1(I, 6, 2),
        repav=fila(I, 8, 2), repav_f1=fila1(I, 9, 2), repav_f2=fila1(I, 10, 2),
        operacion=fila(I, 13, 2), operacion_f1=fila1(I, 14, 2), operacion_f2=fila1(I, 15, 2),
        mantenimiento=fila(I, 17, 2), mantenimiento_f1=fila1(I, 18, 2), mantenimiento_f2=fila1(I, 19, 2),
        servicios=fila(I, 21, 2), servicios_f1=fila1(I, 22, 2), servicios_f2=fila1(I, 23, 2),
        administracion=fila(I, 25, 2), administracion_f1=fila1(I, 26, 2), administracion_f2=fila1(I, 27, 2),
        h_peajes=fila1(H, 18, 3), h_peajes_nuevos=fila1(H, 19, 3), h_otros=fila1(H, 20, 3),
        h_cym=fila1(H, 24, 3), h_atu=fila1(H, 25, 3), h_adm=fila1(H, 26, 3), h_expl=fila1(H, 27, 3),
        h_pv=fila1(H, 31, 3), h_oo=fila1(H, 32, 3), h_rep=fila1(H, 33, 3),
    )
    return inp


# ---------------------------------------------------------------------------
# Funciones financieras (equivalentes a Excel)
# ---------------------------------------------------------------------------
def npv(tasa, flujos):
    """NPV de Excel: el primer flujo se descuenta un período."""
    return float(sum(f / (1 + tasa) ** (t + 1) for t, f in enumerate(flujos)))


def irr(flujos, guess=0.1):
    """TIR. Primero intenta Newton desde 10% (como Excel); si no converge,
    busca un cambio de signo del VAN. Devuelve None si no hay TIR."""
    f = np.asarray(flujos, dtype=float)
    if not (np.any(f > 0) and np.any(f < 0)):
        return None
    t = np.arange(len(f))

    def v(r):
        return float(np.sum(f / (1 + r) ** t))

    def dv(r):
        return float(np.sum(-t * f / (1 + r) ** (t + 1)))

    r = guess
    for _ in range(100):
        try:
            y, d = v(r), dv(r)
        except (OverflowError, ZeroDivisionError, FloatingPointError):
            break
        if d == 0 or not np.isfinite(y) or not np.isfinite(d):
            break
        nr = r - y / d
        if nr <= -0.9999:
            nr = (r - 0.9999) / 2
        if abs(nr - r) < 1e-12:
            r = nr
            if abs(v(r)) < 1e-6 * max(1.0, np.abs(f).max()):
                return r
            break
        r = nr
    if np.isfinite(r) and r > -0.9999 and abs(v(r)) < 1e-6 * max(1.0, np.abs(f).max()):
        return r
    # búsqueda por bisección
    grid = np.concatenate([np.linspace(-0.95, 1, 400), np.linspace(1, 50, 400)])
    vals = [v(x) for x in grid]
    for a, b, va, vb in zip(grid[:-1], grid[1:], vals[:-1], vals[1:]):
        if va == 0:
            return float(a)
        if va * vb < 0:
            for _ in range(200):
                m = (a + b) / 2
                vm = v(m)
                if va * vm <= 0:
                    b, vb = m, vm
                else:
                    a, va = m, vm
            return float((a + b) / 2)
    return None


def mirr(flujos, tasa_fin, tasa_reinv):
    f = np.asarray(flujos, dtype=float)
    n = len(f)
    pos = np.where(f > 0, f, 0.0)
    neg = np.where(f < 0, f, 0.0)
    npv_pos = npv(tasa_reinv, pos)
    npv_neg = npv(tasa_fin, neg)
    if npv_neg == 0 or npv_pos == 0:
        return None
    return float((-npv_pos * (1 + tasa_reinv) ** n / (npv_neg * (1 + tasa_fin))) ** (1 / (n - 1)) - 1)


def pmt(tasa, n, pv):
    if tasa == 0:
        return pv / n
    return pv * tasa / (1 - (1 + tasa) ** (-n))


def round_excel(x, digitos):
    """ROUND de Excel (redondeo 'half away from zero')."""
    factor = 10.0 ** digitos
    return math.copysign(math.floor(abs(x) * factor + 0.5) / factor, x)


def payback(acumulado, plazo):
    """Año desde el cual el flujo acumulado queda positivo (fórmula de FLUJO C60).
    Devuelve un entero o 'No recupera'."""
    negativos = [i + 1 for i, a in enumerate(acumulado) if a < 0]
    if not negativos:
        return 1
    ultimo = negativos[-1]
    return "No recupera" if ultimo >= plazo else ultimo + 1


# ---------------------------------------------------------------------------
# Resultados
# ---------------------------------------------------------------------------
@dataclass
class Resultados:
    kpis: dict
    flujo: pd.DataFrame
    impuestos: pd.DataFrame
    amortizaciones: pd.DataFrame
    prestamo: pd.DataFrame
    capex_opex: pd.DataFrame
    avisos: list


def _amortizacion(base, periodos):
    """Amortiza cada inversión del año s en `periodos[s]` cuotas iguales."""
    out = np.zeros(N)
    for s in range(N):
        if base[s] == 0:
            continue
        n = int(periodos[s])
        for t in range(s, min(s + n, N)):
            out[t] += base[s] / n
    return out


def calcular(inp: Inputs, _sin_deuda_interno: bool = False) -> Resultados:
    P = inp.plazo
    anio = np.arange(1, N + 1)
    activo = (anio <= P).astype(float)          # filas con IF(CONTROL!C12=2, ..., 0)
    g = inp.crecimiento
    a = lambda lst: np.asarray(lst, dtype=float)

    # ---------------- INPUTS / CAPEX / OPEX ----------------
    obras_in = a(inp.obras) * a(inp.obras_f1) * a(inp.obras_f2)                 # INPUTS fila 7
    repav_in = a(inp.repav) * a(inp.repav_f1) * a(inp.repav_f2)                 # INPUTS fila 11
    oper_in = a(inp.operacion) * a(inp.operacion_f1) * a(inp.operacion_f2)      # fila 16
    mant_in = a(inp.mantenimiento) * a(inp.mantenimiento_f1) * a(inp.mantenimiento_f2)  # fila 20
    serv_in = a(inp.servicios) * a(inp.servicios_f1) * a(inp.servicios_f2)      # fila 24
    adm_in = a(inp.administracion) * a(inp.administracion_f1) * a(inp.administracion_f2)  # fila 28

    pv = np.zeros(N); pv[0] = obras_in[0] * (1 + inp.sens_obras)               # CAPEX fila 17 (año 1)
    oo = np.zeros(N); oo[1:] = obras_in[1:] * (1 + inp.sens_obras)             # CAPEX fila 18 (años 2+)
    rep = repav_in * (1 + inp.sens_repav)                                       # CAPEX fila 19

    f_pv = pv * a(inp.h_pv)        # FLUJO 94
    f_oo = oo * a(inp.h_oo)        # FLUJO 95
    f_rep = rep * a(inp.h_rep)     # FLUJO 96
    capex = (f_pv + f_oo + f_rep) * activo   # FLUJO 97 (con IF de escenario)

    cym = mant_in * (1 + inp.sens_cym) * a(inp.h_cym)    # FLUJO 80
    atu = serv_in * (1 + inp.sens_atu) * a(inp.h_atu)    # FLUJO 81
    adm = adm_in * (1 + inp.sens_adm) * a(inp.h_adm)     # FLUJO 82
    expl = oper_in * (1 + inp.sens_expl) * a(inp.h_expl)  # FLUJO 83

    # ---------------- Ingresos ----------------
    base_uteq = inp.uteq
    peajes = np.zeros(N)
    acum = 1.0
    for t in range(1, N):            # años 2..20
        acum *= (1 + g)
        if t <= 4:                   # años 2..5: recalcula desde la base con su sensibilidad
            base = base_uteq
            if inp.uteq_base_desde_anio3 is not None and t >= 2:
                base = inp.uteq_base_desde_anio3
            peajes[t] = base * inp.tarifa * inp.nuevos_peajes[t - 1] * acum * inp.h_peajes[t]
        else:                        # años 6+: encadenado sobre el año anterior
            peajes[t] = peajes[t - 1] * (1 + g) * inp.h_peajes[t]
    adic = np.zeros(N)
    adic[1] = inp.uteq_adicional * inp.tarifa * (1 + g) * inp.h_peajes_nuevos[1]
    for t in range(2, N):
        adic[t] = adic[t - 1] * (1 + g) * inp.h_peajes_nuevos[t]
    otros = (peajes + adic) * inp.otros_ingresos_pct * a(inp.h_otros)
    otros[0] = 0.0

    capex_total_w97 = float(capex.sum())                 # FLUJO W97
    credito = inp.pct_deuda * capex_total_w97            # CONTROL G24
    ing_credito = np.zeros(N); ing_credito[0] = credito  # FLUJO 76
    total_ing = (peajes + adic + otros + ing_credito) * activo   # FLUJO 77 / 22
    tot_ing_sin_if = peajes + adic + otros + ing_credito

    # ---------------- Préstamo ----------------
    saldo0 = credito * (1 + inp.tna) if inp.gracia else credito  # PRESTAMO D2
    gastos = credito * inp.gastos_prestamo_pct                   # PRESTAMO D9
    n_c = max(int(inp.cuotas), 1)
    filas_p = []
    saldo = saldo0
    cuota_fr = pmt(inp.tna, n_c, saldo0) if saldo0 else 0.0
    for k in range(1, n_c + 1):
        interes = inp.tna * saldo
        if inp.sistema == 1:
            cuota = cuota_fr
            capital = cuota - interes
        else:
            capital = saldo0 / n_c
            cuota = interes + capital
        saldo = saldo - capital
        filas_p.append((k, cuota, interes, capital, saldo))
    cuota_k = {k: c for k, c, *_ in filas_p}
    interes_k = {k: i for k, _, i, *_ in filas_p}

    deuda = np.zeros(N)  # FLUJO 84
    for t in range(1, N + 1):
        k = t if not inp.gracia else t - 1
        val = cuota_k.get(k, 0.0)
        if k == 1:
            val += gastos
        deuda[t - 1] = val
    deuda = deuda * activo

    # ---------------- Garantías ----------------
    obras_g12 = float((f_pv + f_oo).sum())
    repav_g13 = float(f_rep.sum())
    opex_g14 = float((cym + atu + adm + expl).sum())
    F44 = inp.gar_pr_contrato * inp.gar_pr_costo_contrato + inp.gar_pr_obras_1a5 * inp.gar_pr_costo_obras
    F45 = inp.gar_pr_contrato * inp.gar_pr_costo_contrato + inp.gar_pr_obras_6a20 * inp.gar_pr_costo_obras
    F50 = round_excel((obras_g12 + repav_g13) * inp.caucion_pct_obras, -8)
    F51 = round_excel(opex_g14 * inp.caucion_pct_contrato, -8)
    F52 = F51
    F55 = F50 * inp.caucion_costo_obras + F51 * inp.caucion_costo_contrato
    F56 = F50 * inp.caucion_costo_obras + F52 * inp.caucion_costo_contrato
    garantias = np.array([F44 + F55 if t <= 5 else F45 + F56 for t in anio], dtype=float)  # FLUJO 99

    # ---------------- Impuestos ----------------
    iva = inp.iva
    ventas = peajes + adic + otros                                   # SUM(FLUJO C73:C75)
    iibb = ventas / 1.21 * inp.iibb                                  # IMPUESTOS 16
    sellos = np.zeros(N); sellos[:5] = capex_total_w97 / (1 + iva) * inp.sellos / 5   # IMPUESTOS 20
    mun = np.zeros(N)
    mun[0] = ventas[0] / (1 + iva) * inp.municipales                 # IMPUESTOS 22 (año 1)
    mun[1:] = total_ing[1:] / (1 + iva) * inp.municipales            # años 2+: usa FLUJO 77
    idc = total_ing * inp.idc                                        # IMPUESTOS 24 (FLUJO 77)

    iva_cred = (cym * inp.iva_cym / (1 + inp.iva_cym) + atu * inp.iva_atu / (1 + inp.iva_atu)
                + adm * inp.iva_adm / (1 + inp.iva_adm) + expl * inp.iva_expl / (1 + inp.iva_expl)
                + capex * inp.iva_capex / (1 + inp.iva_capex) + garantias * iva / (1 + iva))  # IMPUESTOS 33
    iva_deb = ventas * iva / (1 + iva)                               # IMPUESTOS 35
    pos = iva_deb - iva_cred                                         # IMPUESTOS 37
    iva_acum = np.zeros(N); iva_pago = np.zeros(N)
    prev = 0.0
    for t in range(N):
        cur = pos[t] + prev if (pos[t] + prev) < 0 else pos[t]       # IMPUESTOS 39
        if cur < 0:
            pago = 0.0
        elif prev <= 0 and cur < 0:
            pago = 0.0
        else:
            pago = cur + prev if prev < 0 else cur                   # IMPUESTOS 41
        iva_acum[t] = cur; iva_pago[t] = pago
        prev = cur

    # Amortizaciones (hoja AMORTIZACIONES, versión automática)
    base_obras = (f_pv + f_oo) * inp.pct_amortizable / (1 + inp.iva_capex) * activo
    base_rep = f_rep * inp.pct_amortizable / (1 + inp.iva_capex) * activo
    per = lambda vida: np.array([max(1, min(vida, P - s + 1)) for s in anio])
    am = {
        2: _amortizacion(base_obras, per(inp.vida_util_obras_5)) + _amortizacion(base_rep, per(inp.vida_util_repav)),
        1: _amortizacion(base_obras, np.array([max(1, P - s + 1) for s in anio])) + _amortizacion(base_rep, per(inp.vida_util_repav)),
        3: _amortizacion(base_obras, np.ones(N)) + _amortizacion(base_rep, np.ones(N)),
    }
    amort = am.get(inp.variante_amort, am[3])

    ing_gan = np.zeros(N); ing_gan[1:] = total_ing[1:] / (1 + iva)  # IMPUESTOS 46 (año 1 = 0)
    opex_con_iva = cym + atu + adm + expl
    iva_opex = (cym * inp.iva_cym / (1 + inp.iva_cym) + atu * inp.iva_atu / (1 + inp.iva_atu)
                + adm * inp.iva_adm / (1 + inp.iva_adm) + expl * inp.iva_expl / (1 + inp.iva_expl))
    gastos_gan = opex_con_iva - iva_opex + garantias / (1 + iva)     # IMPUESTOS 50
    imp_ded = iibb + sellos + mun + idc                              # IMPUESTOS 54
    base_ai = ing_gan - gastos_gan - imp_ded                         # IMPUESTOS 56
    tope_int = np.where(base_ai > 0, base_ai * inp.pct_deduccion_intereses, 0.0)  # IMPUESTOS 58
    int_ded = np.zeros(N)                                            # IMPUESTOS 60
    for t in range(1, N):
        k = t   # año t+1 usa el interés de la cuota t
        if k > inp.cuotas:
            int_ded[t] = 0.0
        else:
            int_ded[t] = min(tope_int[t], interes_k.get(k, 0.0))
    base_imp = base_ai - int_ded - amort                            # IMPUESTOS 62
    bi_acum = np.zeros(N); gan = np.zeros(N)
    prev = 0.0
    for t in range(N):
        bi_acum[t] = 0.0 if (base_imp[t] + prev) > 0 else base_imp[t] + prev   # IMPUESTOS 64
        if base_imp[t] < 0 or (base_imp[t] + prev) < 0:
            gan[t] = 0.0
        else:
            gan[t] = (base_imp[t] + prev) * inp.ganancias                       # IMPUESTOS 66
        prev = bi_acum[t]

    impuestos = iva_pago + gan + iibb + mun + sellos + idc           # FLUJO 85:90

    # ---------------- Flujo ----------------
    opex_total = (cym + atu + adm + expl + deuda + impuestos) * activo   # FLUJO 31
    resultado_op = total_ing - opex_total - garantias                    # FLUJO 33
    flujo_neto = resultado_op - capex                                    # FLUJO 45
    egresos = opex_total + capex + garantias                             # FLUJO 43
    acumulado = np.cumsum(flujo_neto)
    acumulado_if = np.where(activo > 0, acumulado, 0.0)                  # FLUJO 46 con IF

    fl_tir = flujo_neto[:P]
    tir = irr(fl_tir)
    tirm = mirr(fl_tir, inp.tna, inp.tasa_descuento)
    van = npv(inp.tasa_descuento, flujo_neto)
    van_egr = npv(inp.tasa_descuento, egresos)
    retorno = van / van_egr if van_egr else None
    vai_vae = npv(inp.tasa_descuento, total_ing) / van_egr if van_egr else None
    pb = payback(acumulado_if, P)
    desc = np.cumsum(flujo_neto / (1 + inp.tasa_descuento) ** anio)
    pb_desc = payback(np.where(activo > 0, desc, 0.0), P)

    kpis = {
        "TIR (con IVA, incluye préstamo)": tir,
        "TIR modificada": tirm,
        "VAN": van,
        "Retorno (VAN / VA egresos)": retorno,
        "VAI / VAE": vai_vae,
        "Payback (años)": pb,
        "Payback descontado (años)": pb_desc,
        "Crédito": credito,
    }
    if not _sin_deuda_interno:
        sd = calcular(inp.copia(pct_deuda=0.0), _sin_deuda_interno=True)
        kpis["TIR del proyecto (sin préstamo)"] = sd.kpis["TIR (con IVA, incluye préstamo)"]
        kpis["VAN del proyecto (sin préstamo)"] = sd.kpis["VAN"]

    cols = [f"Año {i} ({inp.anio_inicio + i - 1})" for i in anio]
    flujo = pd.DataFrame({
        "Peajes": peajes * activo, "Peajes adicionales": adic * activo, "Otros ingresos": otros * activo,
        "Ingreso crédito": ing_credito * activo, "TOTAL INGRESOS": total_ing,
        "Conservación y mantenimiento": cym * activo, "Atención al usuario": atu * activo,
        "Administración": adm * activo, "Explotación": expl * activo,
        "Servicio de la deuda": deuda, "Impuestos": impuestos * activo, "TOTAL OPEX": opex_total,
        "Obras de puesta en valor": f_pv * activo, "Obras obligatorias": f_oo * activo,
        "Repavimentación": f_rep * activo, "TOTAL CAPEX": capex, "Garantías": garantias,
        "TOTAL EGRESOS": egresos, "FLUJO NETO": flujo_neto, "FLUJO NETO ACUMULADO": acumulado_if,
    }, index=cols).T
    imp_df = pd.DataFrame({
        "IVA débito": iva_deb, "IVA crédito": iva_cred, "Posición IVA": pos, "IVA acumulado": iva_acum,
        "IVA a pagar": iva_pago, "Ingresos Brutos": iibb, "Sellos": sellos, "Tasas municipales": mun,
        "Débitos y créditos": idc, "Ingresos gravados (Ganancias)": ing_gan, "Gastos deducibles": gastos_gan,
        "Impuestos deducibles": imp_ded, "Base antes de intereses y amortizaciones": base_ai,
        "Intereses deducibles": int_ded, "Amortización": amort, "Base imponible": base_imp,
        "Quebranto acumulado": bi_acum, "Impuesto a las Ganancias": gan, "TOTAL IMPUESTOS": impuestos,
    }, index=cols).T
    am_df = pd.DataFrame({
        "Base obras (sin IVA)": base_obras, "Base repavimentación (sin IVA)": base_rep,
        "Variante 1: obras plazo restante + repav vida útil": am[1],
        "Variante 2: todo a vida útil 5 años": am[2],
        "Variante 3: instantánea": am[3],
        "Amortización usada": amort,
    }, index=cols).T
    pr_df = pd.DataFrame(filas_p, columns=["Cuota N°", "Cuota", "Interés", "Capital", "Saldo"])
    co_df = pd.DataFrame({
        "Obras (INPUTS)": obras_in, "Repavimentación (INPUTS)": repav_in, "Operación": oper_in,
        "Mantenimiento": mant_in, "Servicios al usuario": serv_in, "Administración": adm_in,
    }, index=cols).T

    avisos = []
    if tir is None:
        avisos.append("La TIR no se puede calcular con este flujo (no tiene un cambio de signo claro).")
    elif tir > 1:
        avisos.append("La TIR supera el 100%: el préstamo cubre casi toda la inversión inicial y la TIR pierde sentido económico. "
                      "Conviene informar la TIR del proyecto sin préstamo.")
    if flujo_neto[0] > 0:
        avisos.append("El flujo del año 1 es positivo porque el crédito supera la inversión de ese año.")
    if pb == "No recupera":
        avisos.append("El proyecto no recupera la inversión dentro del plazo.")

    return Resultados(kpis, flujo, imp_df, am_df, pr_df, co_df, avisos)
